import re
import requests
from bs4 import BeautifulSoup
import openpyxl
from config import SCHEDULE_PAGE_URL, LOCAL_FILE_NAME, TARGET_GROUP, TARGET_SHEET

def get_real_cell_value(sheet, row: int, col: int):
    cell_val = sheet.cell(row=row, column=col).value
    if cell_val is not None:
        return cell_val
    
    for rng in sheet.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return sheet.cell(row=rng.min_row, column=rng.min_col).value
            
    return None

def extract_lesson_details(raw_text: str):
    lt_lower = raw_text.lower()

    lesson_type = "занятие"
    if re.search(r'\b(сем|семинар|сем\.)\b', lt_lower):
        lesson_type = "семинар"
    elif re.search(r'\b(прак|практика|лаб|лабораторная)\b', lt_lower):
        lesson_type = "практика"
    elif re.search(r'\b(лек|лекция|лек\.)\b', lt_lower):
        lesson_type = "лекция"

    room = "не указана"
    if "ауд" in lt_lower:
        parts = lt_lower.split("ауд")
        if len(parts) > 1:
            r_cand = parts[1].strip(" .:")
            if r_cand:
                room = f"ауд. {r_cand.split()[0]}"
    elif "онлайн" in lt_lower:
        room = "Онлайн"

    comma_parts = [p.strip() for p in raw_text.split(",") if p.strip()]
    subject = comma_parts[0] if len(comma_parts) > 0 else raw_text
    
    teacher = "не указан"
    if len(comma_parts) > 2:
        t_part = comma_parts[2]
        if "ауд" in t_part.lower():
            t_part = t_part[:t_part.lower().find("ауд")].strip()
        teacher = t_part if t_part else "не указан"
    elif len(comma_parts) == 2 and not re.search(r'\b(лек|сем|прак|лаб)\b', comma_parts[1].lower()):
        teacher = comma_parts[1]

    return subject, lesson_type, teacher, room

import urllib.parse
import requests
from bs4 import BeautifulSoup

# Прямой fallback-URL на случай динамического рендеринга страницы
FALLBACK_SCHEDULE_URL = "https://esil.edu.kz/api/media/file/%2B%D0%A0%D0%B0%D1%81%D0%BF%D0%B8%D1%81%D0%B0%D0%BD%D0%B8%D0%B5_2%D0%BA-4%D0%B3%20%2B1%D0%BA3%D0%B3%2B1%D0%BA2%D0%B3%20009-6.xlsx"

def download_schedule_file():
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    }
    target_url = None

    try:
        res = requests.get(SCHEDULE_PAGE_URL, headers=headers, timeout=15)
        if res.status_code == 200:
            soup = BeautifulSoup(res.text, 'html.parser')
            
            # Поиск ссылки по атрибутам href и декодированным именам файлов
            for link in soup.find_all('a', href=True):
                href = link['href'].strip()
                decoded_href = urllib.parse.unquote(href).lower()
                
                if ("/api/media/file/" in href or ".xlsx" in decoded_href) and any(k in decoded_href for k in ["2к", "2_к", "2-к", "2 курс"]):
                    target_url = href
                    break
    except Exception as e:
        print(f"Ошибка при парсинге страницы: {e}")

    # Если сайт использует pure JS рендеринг и ссылка не найдена в static HTML
    if not target_url:
        target_url = FALLBACK_SCHEDULE_URL

    if not target_url.startswith("http"):
        target_url = "https://esil.edu.kz" + target_url

    print(f"Загрузка файла расписания: {target_url}")
    f_res = requests.get(target_url, headers=headers, timeout=30)
    f_res.raise_for_status()

    with open(LOCAL_FILE_NAME, "wb") as f:
        f.write(f_res.content)

def parse_excel_schedule():
    wb = openpyxl.load_workbook(LOCAL_FILE_NAME, data_only=True)
    if TARGET_SHEET not in wb.sheetnames:
        raise Exception(f"Лист '{TARGET_SHEET}' не найден в файле.")
    sheet = wb[TARGET_SHEET]

    group_col = None
    for c in range(1, sheet.max_column + 1):
        val = sheet.cell(row=14, column=c).value
        if val and TARGET_GROUP in str(val):
            group_col = c
            break

    if group_col is None:
        raise Exception(f"Колонка для группы {TARGET_GROUP} не найдена в строке 14.")

    day_mapping = {
        "понедельник": 0, "вторник": 1, "среда": 2, 
        "четверг": 3, "пятница": 4, "суббота": 5, "воскресенье": 6
    }

    parsed_lessons = []
    current_day = None
    current_time = None

    for r in range(15, sheet.max_row + 1):
        d_val = get_real_cell_value(sheet, r, 1)
        if d_val:
            d_str = str(d_val).strip().lower()
            for d_name, d_code in day_mapping.items():
                if d_name in d_str:
                    current_day = d_code
                    break

        t_val = get_real_cell_value(sheet, r, 2)
        if t_val:
            t_str = str(t_val).strip()
            if "-" in t_str or "–" in t_str:
                current_time = t_str.replace("–", "-")

        raw_val = get_real_cell_value(sheet, r, group_col)

        if raw_val and current_day is not None and current_time:
            raw_text = str(raw_val).strip()
            if not raw_text or raw_text.lower() == "nan":
                continue

            time_clean = current_time.replace(".", ":")
            times = time_clean.split("-")
            t_start = times[0].strip() if len(times) > 0 else ""
            t_end = times[1].strip() if len(times) > 1 else ""

            subject, lesson_type, teacher, room = extract_lesson_details(raw_text)

            item = (current_day, t_start, t_end, subject, lesson_type, teacher, room)
            if item not in parsed_lessons:
                parsed_lessons.append(item)

    return parsed_lessons