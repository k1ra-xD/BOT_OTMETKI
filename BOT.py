import math
import asyncio
import os
import requests
from bs4 import BeautifulSoup
import openpyxl
from datetime import datetime, timedelta, timezone
import asyncpg
from aiohttp import web
from aiogram import Bot, Dispatcher, F, Router, BaseMiddleware
from aiogram.filters import Command
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import (
    Message, CallbackQuery, 
    ReplyKeyboardMarkup, KeyboardButton,
    InlineKeyboardMarkup, InlineKeyboardButton,
    ReplyKeyboardRemove, WebAppInfo
)

# --- КОНФИГУРАЦИЯ ---
TOKEN = os.environ.get("BOT_TOKEN", "8932791447:AAGB5HfDMv1Jq7yMwVwko9YVl7rubu7F3tM")
DATABASE_URL = os.environ.get("DATABASE_URL", "postgresql://postgres.stduhsqtkysawtahgftq:7htTw1GOgUV7XXb2@aws-0-eu-west-1.pooler.supabase.com:6543/postgres")
ADMIN_ID = 1231388093  # Ваш ID

SCHEDULE_PAGE_URL = "https://esil.edu.kz/students_schedule/"
LOCAL_FILE_NAME = "target_schedule.xlsx"
TARGET_GROUP = "Б-ИТЗД 26/23 Р" 
TARGET_SHEET = "рус Прик 1,3 1,2"

DEFAULT_UNI_LAT = 51.159555 
DEFAULT_UNI_LON = 71.458555
DEFAULT_RADIUS = 150

ASTANA_TZ = timezone(timedelta(hours=5))

db_pool: asyncpg.Pool = None
router = Router()

DAYS_MAP = {
    0: "Понедельник",
    1: "Вторник",
    2: "Среда",
    3: "Четверг",
    4: "Пятница",
    5: "Суббота",
    6: "Воскресенье"
}

# --- ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ ПАРСИНГА EXCEL ---
def get_real_cell_value(sheet, row: int, col: int):
    """
    Возвращает реальное значение ячейки с учетом горизонтального/вертикального объединения (Merged Cells).
    """
    cell_val = sheet.cell(row=row, column=col).value
    if cell_val is not None:
        return cell_val
    
    for rng in sheet.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return sheet.cell(row=rng.min_row, column=rng.min_col).value
            
    return None

def extract_lesson_details(raw_text: str):
    """ Разбирает сырой текст ячейки на тип занятия, преподавателя и аудиторию. """
    lesson_type = "занятие"
    lt_lower = raw_text.lower()
    if "лек" in lt_lower:
        lesson_type = "лекция"
    elif "сем" in lt_lower:
        lesson_type = "семинар"
    elif "прак" in lt_lower:
        lesson_type = "практика"

    room = "не указана"
    if "ауд" in lt_lower:
        parts = lt_lower.split("ауд")
        if len(parts) > 1:
            r_cand = parts[1].strip(" .:")
            if r_cand:
                room = f"ауд. {r_cand.split()[0]}"
    elif "онлайн" in lt_lower:
        room = "Онлайн"

    comma_parts = [p.strip() for p in raw_text.split(",") if p.strip()]
    subject = comma_parts[0] if len(comma_parts) > 0 else raw_text
    
    teacher = "не указан"
    if len(comma_parts) > 2:
        t_part = comma_parts[2]
        if "ауд" in t_part.lower():
            t_part = t_part[:t_part.lower().find("ауд")].strip()
        teacher = t_part if t_part else "не указан"
    elif len(comma_parts) == 2 and not any(k in comma_parts[1].lower() for k in ["лек", "сем", "прак"]):
        teacher = comma_parts[1]

    return subject, lesson_type, teacher, room


# --- ИНИЦИАЛИЗАЦИЯ БД ---
async def init_db():
    async with db_pool.acquire() as conn:
        await conn.execute("""
            CREATE TABLE IF NOT EXISTS students (
                telegram_id BIGINT PRIMARY KEY,
                full_name TEXT
            );
            CREATE TABLE IF NOT EXISTS banned_users (
                telegram_id BIGINT PRIMARY KEY
            );
            CREATE TABLE IF NOT EXISTS settings (
                key TEXT PRIMARY KEY,
                value TEXT
            );
            CREATE TABLE IF NOT EXISTS schedule (
                id SERIAL PRIMARY KEY,
                day_of_week INT,
                time_start TEXT,
                time_end TEXT,
                subject TEXT,
                lesson_type TEXT,
                teacher TEXT,
                room TEXT
            );
        """)
        
        await conn.execute("INSERT INTO settings (key, value) VALUES ('lat', $1) ON CONFLICT (key) DO NOTHING", str(DEFAULT_UNI_LAT))
        await conn.execute("INSERT INTO settings (key, value) VALUES ('lon', $1) ON CONFLICT (key) DO NOTHING", str(DEFAULT_UNI_LON))
        await conn.execute("INSERT INTO settings (key, value) VALUES ('radius', $1) ON CONFLICT (key) DO NOTHING", str(DEFAULT_RADIUS))


async def get_setting(key: str, default):
    async with db_pool.acquire() as conn:
        row = await conn.fetchrow("SELECT value FROM settings WHERE key = $1", key)
        if not row:
            return default
        val = row['value']
        if key in ['lat', 'lon']:
            return float(val)
        elif key == 'radius':
            return int(val)
        return val

async def set_setting(key: str, value):
    async with db_pool.acquire() as conn:
        await conn.execute("""
            INSERT INTO settings (key, value) VALUES ($1, $2)
            ON CONFLICT (key) DO UPDATE SET value = EXCLUDED.value
        """, key, str(value))


# --- Middleware для проверки блокировки ---
class BannedMiddleware(BaseMiddleware):
    async def __call__(self, handler, event: Message, data):
        user = data.get("event_from_user")
        if user and user.id != ADMIN_ID:
            async with db_pool.acquire() as conn:
                banned = await conn.fetchrow("SELECT telegram_id FROM banned_users WHERE telegram_id = $1", user.id)
            
            if banned:
                if isinstance(event, Message):
                    await event.answer("⛔️ Доступ закрыт. Вы заблокированы администратором.", reply_markup=ReplyKeyboardRemove())
                return
        return await handler(event, data)

router.message.middleware(BannedMiddleware())


# Состояния FSM
class RegStates(StatesGroup):
    waiting_for_name = State()
    waiting_for_radius = State()
    waiting_for_ban_id = State()
    waiting_for_unban_id = State()
    waiting_for_uni_location = State()

class ScheduleAdminStates(StatesGroup):
    waiting_for_schedule_text = State()


current_session = {
    "is_active": False,
    "responses": {}
}


def calculate_distance(lat1, lon1, lat2, lon2):
    R = 6371000   
    phi1, phi2 = math.radians(lat1), math.radians(lat2)
    dphi = math.radians(lat2 - lat1)
    dlambda = math.radians(lon2 - lon1)

    a = math.sin(dphi/2)**2 + math.cos(phi1) * math.cos(phi2) * math.sin(dlambda/2)**2
    c = 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))
    return R * c


# --- ФОНОВЫЙ ТАЙМЕР НАПОМИНАНИЙ ЗА 10 МИНУТ ---
notified_lessons = set()

async def schedule_notifications_loop(bot: Bot):
    while True:
        try:
            now = datetime.now(ASTANA_TZ)
            current_day = now.weekday()
            target_time = (now + timedelta(minutes=10)).strftime("%H:%M")
            date_key = now.strftime("%Y-%m-%d")

            async with db_pool.acquire() as conn:
                lessons = await conn.fetch("""
                    SELECT * FROM schedule 
                    WHERE day_of_week = $1 AND time_start = $2
                """, current_day, target_time)

                if lessons:
                    students = await conn.fetch("SELECT telegram_id FROM students")
                    for lesson in lessons:
                        lesson_id_key = f"{date_key}_{lesson['id']}"
                        if lesson_id_key not in notified_lessons:
                            notified_lessons.add(lesson_id_key)
                            
                            text = (
                                f"🔔 <b>Напоминание о паре!</b>\n"
                                f"⏰ <b>Начало:</b> через 10 минут ({lesson['time_start']} - {lesson['time_end']})\n"
                                f"📖 <b>Предмет:</b> {lesson['subject']} ({lesson['lesson_type']})\n"
                                f"👨‍🏫 <b>Преподаватель:</b> {lesson['teacher']}\n"
                                f"🚪 <b>Аудитория:</b> {lesson['room']}"
                            )
                            
                            for student in students:
                                try:
                                    await bot.send_message(student['telegram_id'], text, parse_mode="HTML")
                                except Exception:
                                    pass
        except Exception as e:
            print(f"Ошибка в рассылке расписания: {e}")

        await asyncio.sleep(30)


# --- КЛАВИАТУРЫ ---
student_main_kb = ReplyKeyboardMarkup(
    keyboard=[
        [KeyboardButton(text="📅 Расписание"), KeyboardButton(text="📍 Отметиться")]
    ],
    resize_keyboard=True
)

schedule_inline_kb = InlineKeyboardMarkup(inline_keyboard=[
    [
        InlineKeyboardButton(text="Сегодня", callback_data="sched_today"),
        InlineKeyboardButton(text="Завтра", callback_data="sched_tomorrow")
    ],
    [InlineKeyboardButton(text="📅 На всю неделю", callback_data="sched_week")]
])


# --- ПАНЕЛЬ УПРАВЛЕНИЯ И СТАРТ ---
@router.message(Command("start"))
async def cmd_start(message: Message, state: FSMContext):
    if message.from_user.id == ADMIN_ID:
        radius = await get_setting('radius', DEFAULT_RADIUS)
        lat = await get_setting('lat', DEFAULT_UNI_LAT)
        lon = await get_setting('lon', DEFAULT_UNI_LON)
        
        admin_kb = ReplyKeyboardMarkup(
            keyboard=[
                [KeyboardButton(text="🟢 Проверить присутствие"), KeyboardButton(text="🔴 Завершить проверку")],
                [KeyboardButton(text="👀 Кто ответил"), KeyboardButton(text="🎯 Установить центр ВУЗа")],
                [KeyboardButton(text="👥 Список студентов"), KeyboardButton(text="🗑 Удалить студента")],
                [KeyboardButton(text="🚫 Блокировать ID"), KeyboardButton(text="✅ Разблокировать ID")],
                [KeyboardButton(text="⚙️ Изменить радиус зоны"), KeyboardButton(text="📅 Расписание")]
            ],
            resize_keyboard=True
        )

        admin_inline_panel = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="📊 WebApp Дашборд", web_app=WebAppInfo(url="https://bot-otmetki.onrender.com/dashboard"))],
            [InlineKeyboardButton(text="🔄 Авто-обновить с сайта (esil.edu.kz)", callback_data="admin_auto_update_schedule")],
            [InlineKeyboardButton(text="✏️ Ввести расписание вручную", callback_data="admin_edit_schedule")]
        ])

        await message.answer(
            f"👋 <b>Панель администратора</b>\n\n"
            f"📍 Координаты ВУЗа: <code>{lat}, {lon}</code>\n"
            f"📏 Радиус зоны: <b>{radius}м</b>",
            reply_markup=admin_kb,
            parse_mode="HTML"
        )
        await message.answer("Управление расписанием и статистикой:", reply_markup=admin_inline_panel)
        return

    async with db_pool.acquire() as conn:
        row = await conn.fetchrow("SELECT full_name FROM students WHERE telegram_id = $1", message.from_user.id)

    if row:
        await message.answer(
            f"Привет, {row['full_name']}! Вы зарегистрированы в системе.",
            reply_markup=student_main_kb
        )
    else:
        await message.answer("Привет! Для регистрации отправьте свои **Имя и Фамилию** (например: *Иван Иванов*).")
        await state.set_state(RegStates.waiting_for_name)

@router.message(RegStates.waiting_for_name)
async def process_name(message: Message, state: FSMContext):
    full_name = message.text.strip()
    if len(full_name.split()) < 2:
        await message.answer("Пожалуйста, введите корректно Имя и Фамилию.")
        return

    async with db_pool.acquire() as conn:
        await conn.execute("""
            INSERT INTO students (telegram_id, full_name) VALUES ($1, $2)
            ON CONFLICT (telegram_id) DO UPDATE SET full_name = EXCLUDED.full_name
        """, message.from_user.id, full_name)

    await state.clear()
    await message.answer(f"Спасибо, {full_name}! Регистрация прошла успешно.", reply_markup=student_main_kb)


# --- ХЭНДЛЕРЫ ПРОСМОТРА РАСПИСАНИЯ ---
@router.message(F.text == "📅 Расписание")
async def show_schedule_menu(message: Message):
    await message.answer("Выберите период для просмотра расписания:", reply_markup=schedule_inline_kb)

async def format_day_schedule(conn, day_code: int) -> str:
    lessons = await conn.fetch("""
        SELECT * FROM schedule WHERE day_of_week = $1 ORDER BY time_start ASC
    """, day_code)
    
    day_name = DAYS_MAP.get(day_code, "")
    if not lessons:
        return f"<b>{day_name}:</b>\n🎉 Парочек нет, отдыхаем!\n"
    
    text = f"🗓 <b>{day_name}:</b>\n\n"
    for l in lessons:
        teacher_str = f"👨‍🏫 {l['teacher']} | " if l['teacher'] != 'не указан' else ""
        text += (
            f"⏰ <code>{l['time_start']} - {l['time_end']}</code>\n"
            f"📖 <b>{l['subject']}</b> ({l['lesson_type']})\n"
            f"{teacher_str}🚪 <b>{l['room']}</b>\n"
            f"-------------------------------\n"
        )
    return text

@router.callback_query(F.data.startswith("sched_"))
async def process_schedule_callback(callback: CallbackQuery):
    now = datetime.now(ASTANA_TZ)
    async with db_pool.acquire() as conn:
        if callback.data == "sched_today":
            day_code = now.weekday()
            text = await format_day_schedule(conn, day_code)
        elif callback.data == "sched_tomorrow":
            day_code = (now.weekday() + 1) % 7
            text = await format_day_schedule(conn, day_code)
        elif callback.data == "sched_week":
            text = "📅 <b>Расписание на всю неделю:</b>\n\n"
            for d in range(5):
                text += await format_day_schedule(conn, d) + "\n"
        
        await callback.message.edit_text(text, parse_mode="HTML", reply_markup=schedule_inline_kb)
    await callback.answer()


# --- АВТОМАТИЧЕСКОЕ И РУЧНОЕ ОБНОВЛЕНИЕ РАСПИСАНИЯ ---
async def update_schedule_workflow(status_msg: Message):
    """ Фоновая задача с поэтапным информированием админа о прогрессе парсинга. """
    headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"}
    
    # ЭТАП 1: Скачивание файла с сайта
    try:
        await status_msg.edit_text("⏳ <b>[1/5]</b> Подключение к <code>esil.edu.kz</code> и поиск файла...", parse_mode="HTML")
        
        def download_file():
            res = requests.get(SCHEDULE_PAGE_URL, headers=headers, timeout=30)
            res.raise_for_status()
            soup = BeautifulSoup(res.text, 'html.parser')
            
            target_url = None
            for row in soup.find_all('tr'):
                text = row.get_text()
                if "Расписание занятий 2 курса (4 года)" in text or "1 курса (3 года)" in text:
                    link_tag = row.find('a', href=True)
                    if link_tag:
                        href = link_tag['href']
                        target_url = href if href.startswith("http") else "https://esil.edu.kz" + href
                        break
            if not target_url:
                raise Exception("Не найдена ссылка на файл расписания 2 курса на сайте.")
                
            f_res = requests.get(target_url, headers=headers, timeout=30)
            f_res.raise_for_status()
            with open(LOCAL_FILE_NAME, "wb") as f:
                f.write(f_res.content)

        await asyncio.to_thread(download_file)
        await status_msg.edit_text("✅ <b>[1/5]</b> Файл расписания успешно скачан!\n⏳ <b>[2/5]</b> Открытие структуры Excel...", parse_mode="HTML")
    except Exception as e:
        await status_msg.edit_text(f"❌ <b>Ошибка на этапе 1 (скачивание):</b> {e}", parse_mode="HTML")
        return

    # ЭТАП 2, 3 и 4: Открытие, поиск группы и разбор с учетом объединений
    try:
        def parse_excel():
            wb = openpyxl.load_workbook(LOCAL_FILE_NAME, data_only=True)
            if TARGET_SHEET not in wb.sheetnames:
                raise Exception(f"Лист '{TARGET_SHEET}' не найден в файле.")
            sheet = wb[TARGET_SHEET]

            group_col = None
            for c in range(1, sheet.max_column + 1):
                val = sheet.cell(row=14, column=c).value
                if val and TARGET_GROUP in str(val):
                    group_col = c
                    break

            if group_col is None:
                raise Exception(f"Колонка для группы {TARGET_GROUP} не найдена в строке 14.")

            day_mapping = {
                "понедельник": 0, "вторник": 1, "среда": 2, 
                "четверг": 3, "пятница": 4, "суббота": 5, "воскресенье": 6
            }

            parsed_lessons = []
            current_day = None
            current_time = None

            for r in range(15, sheet.max_row + 1):
                d_val = get_real_cell_value(sheet, r, 1)
                if d_val:
                    d_str = str(d_val).strip().lower()
                    for d_name, d_code in day_mapping.items():
                        if d_name in d_str:
                            current_day = d_code
                            break

                t_val = get_real_cell_value(sheet, r, 2)
                if t_val:
                    t_str = str(t_val).strip()
                    if "-" in t_str or "–" in t_str:
                        current_time = t_str.replace("–", "-")

                raw_val = get_real_cell_value(sheet, r, group_col)

                if raw_val and current_day is not None and current_time:
                    raw_text = str(raw_val).strip()
                    if not raw_text or raw_text.lower() == "nan":
                        continue

                    time_clean = current_time.replace(".", ":")
                    times = time_clean.split("-")
                    t_start = times[0].strip() if len(times) > 0 else ""
                    t_end = times[1].strip() if len(times) > 1 else ""

                    subject, lesson_type, teacher, room = extract_lesson_details(raw_text)

                    item = (current_day, t_start, t_end, subject, lesson_type, teacher, room)
                    if item not in parsed_lessons:
                        parsed_lessons.append(item)

            return group_col, parsed_lessons

        await status_msg.edit_text("⏳ <b>[3/5]</b> Поиск группы Б-ИТЗД 26/23 Р и разрешение объединенных ячеек...", parse_mode="HTML")
        group_col_idx, parsed_lessons = await asyncio.to_thread(parse_excel)
        
        await status_msg.edit_text(
            f"✅ <b>[4/5]</b> Извлечение предметов завершено! Найдено